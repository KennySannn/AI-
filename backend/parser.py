"""
parser.py —— 资料文本提取与分块
支持格式（需求文档 2.1/2.2）：pdf / ppt / pptx / doc / docx / md / txt
说明：不使用向量模型，提取纯文本后按固定长度分块入库。
"""
import os

CHUNK_SIZE = 800  # 每块目标字符数


def _read_plain_text(file_path):
    """md/txt：依次尝试 utf-8 / gbk 编码"""
    for encoding in ('utf-8', 'gbk'):
        try:
            with open(file_path, 'r', encoding=encoding) as f:
                return f.read()
        except UnicodeDecodeError:
            continue
    raise ValueError('无法识别文件编码（已尝试 utf-8 / gbk）')


def _extract_pdf(file_path):
    from pypdf import PdfReader
    reader = PdfReader(file_path)
    pages = [(page.extract_text() or '') for page in reader.pages]
    return '\n'.join(pages)


def _extract_pptx(file_path):
    from pptx import Presentation
    prs = Presentation(file_path)
    parts = []
    for slide in prs.slides:
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()
                if t:
                    texts.append(t)
        if texts:
            parts.append('\n'.join(texts))
    return '\n'.join(parts)


def _extract_docx(file_path):
    from docx import Document
    doc = Document(file_path)
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    # 表格里的文字也要提取
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(' | '.join(cells))
    return '\n'.join(parts)


_EXTRACTORS = {
    'pdf': _extract_pdf,
    'pptx': _extract_pptx,
    'ppt': _extract_pptx,     # 旧版 .ppt 由 python-pptx 尝试，失败则报错
    'docx': _extract_docx,
    'doc': _extract_docx,     # 旧版 .doc 由 python-docx 尝试，失败则报错
    'md': _read_plain_text,
    'txt': _read_plain_text,
}


def extract_text(file_path, file_type):
    """按文件类型提取全部文本。失败抛异常（由调用方标记 status=failed）。"""
    if not os.path.exists(file_path):
        raise ValueError('文件不存在: %s' % file_path)
    extractor = _EXTRACTORS.get(file_type)
    if extractor is None:
        raise ValueError('不支持的文件类型: %s' % file_type)
    text = extractor(file_path)
    if not text or not text.strip():
        raise ValueError('未从文件中提取到文本内容')
    return text


def chunk_text(text, size=CHUNK_SIZE):
    """分块：按段落（空行/换行）聚合到目标长度，超长段落硬切。"""
    text = text.replace('\r\n', '\n').strip()
    paragraphs = [p.strip() for p in text.split('\n') if p.strip()]

    # 把超长段落先硬切成不超过 size 的片段
    pieces = []
    for p in paragraphs:
        while len(p) > size:
            pieces.append(p[:size])
            p = p[size:]
        if p:
            pieces.append(p)

    # 相邻片段聚合到目标长度
    chunks = []
    buf = ''
    for piece in pieces:
        if buf and len(buf) + len(piece) + 1 > size:
            chunks.append(buf)
            buf = piece
        else:
            buf = buf + '\n' + piece if buf else piece
    if buf:
        chunks.append(buf)
    return chunks
